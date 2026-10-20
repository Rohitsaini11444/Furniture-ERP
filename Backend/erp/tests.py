import io
import openpyxl
from decimal import Decimal
import uuid
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase

from erp.models import (
    User, ProductionUnit, StoreItem, StoreItemCategory, StoreItemStatus,
    StoreDailyIssue, StoreMaterialReturn, StoreMaterialIn, Supplier, RoleChoices
)


class MonthlyContractorBillingTests(APITestCase):
    """
    Focused test suite for MonthlyContractorBillingView.
    Validates factory-specific isolation, parameter requirements,
    decimal accuracy, return-credit exclusion, and unassigned transaction detection.
    """

    def setUp(self):
        # Admin user for authentication
        self.admin = User.objects.create_user(
            username='admin_billing_test',
            password='testpassword123',
            role=RoleChoices.ADMIN
        )
        self.client.force_authenticate(user=self.admin)

        # Production units (factories)
        self.unit4 = ProductionUnit.objects.create(
            name='Unit 4',
            unit_code='U4_TEST'
        )
        self.unit2 = ProductionUnit.objects.create(
            name='Unit 2',
            unit_code='U2_TEST'
        )

        # Contractor
        self.contractor = User.objects.create_user(
            username='contractor_ramesh',
            first_name='Ramesh',
            last_name='Carpenter',
            role=RoleChoices.CONTRACTOR,
            production_unit=self.unit4
        )

        # Store items
        self.cat = StoreItemCategory.objects.create(name='Timber & Board')
        self.item1 = StoreItem.objects.create(
            item_code='IT-PLY-01',
            item_name='Marine Plywood 18mm',
            category=self.cat,
            unit='pcs',
            base_rate=Decimal('150.00'),
            current_rate=Decimal('150.00')
        )
        self.item2 = StoreItem.objects.create(
            item_code='IT-FVI-02',
            item_name='Fevicol SR 505',
            category=self.cat,
            unit='ltr',
            base_rate=Decimal('200.00'),
            current_rate=Decimal('200.00')
        )

        self.bill_url = reverse('store-contractor-bill')

    def test_missing_required_parameters(self):
        """API rejects missing contractor, production_unit, or month."""
        # 1. Missing all
        res = self.client.get(self.bill_url)
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('contractor parameter is required', res.data.get('error', ''))

        # 2. Missing production_unit
        res = self.client.get(self.bill_url, {'contractor': str(self.contractor.id)})
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('production_unit parameter is required', res.data.get('error', ''))

        # 3. Missing month
        res = self.client.get(self.bill_url, {
            'contractor': str(self.contractor.id),
            'production_unit': str(self.unit4.id)
        })
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('month parameter is required', res.data.get('error', ''))

    def test_invalid_contractor_or_unit_not_found(self):
        """API returns 404 for invalid UUID or non-existent entity."""
        fake_uuid = str(uuid.uuid4())
        # Invalid contractor
        res = self.client.get(self.bill_url, {
            'contractor': fake_uuid,
            'production_unit': str(self.unit4.id),
            'month': 'Jul-26'
        })
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)

        # Invalid unit
        res = self.client.get(self.bill_url, {
            'contractor': str(self.contractor.id),
            'production_unit': fake_uuid,
            'month': 'Jul-26'
        })
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)

    def test_unit4_billing_excludes_unit2_issues(self):
        """Factory-specific: Unit 4 bill must exclude Unit 2 issues and vice versa."""
        # Issue in Unit 4: 10 pcs @ 150.00 = 1,500.00 (Chargeable)
        StoreDailyIssue.objects.create(
            voucher_no='VCH-U4-001',
            contractor=self.contractor,
            item=self.item1,
            qty=Decimal('10.000'),
            unit='pcs',
            rate=Decimal('150.00'),
            status=StoreItemStatus.CHARGE,
            production_unit=self.unit4,
            month_year='Jul-26',
            issued_by=self.admin
        )

        # Issue in Unit 2: 5 pcs @ 200.00 = 1,000.00 (Chargeable)
        StoreDailyIssue.objects.create(
            voucher_no='VCH-U2-001',
            contractor=self.contractor,
            item=self.item2,
            qty=Decimal('5.000'),
            unit='ltr',
            rate=Decimal('200.00'),
            status=StoreItemStatus.CHARGE,
            production_unit=self.unit2,
            month_year='Jul-26',
            issued_by=self.admin
        )

        # Query Unit 4 statement
        res_u4 = self.client.get(self.bill_url, {
            'contractor': str(self.contractor.id),
            'production_unit': str(self.unit4.id),
            'month': 'Jul-26'
        })
        self.assertEqual(res_u4.status_code, status.HTTP_200_OK)
        self.assertEqual(res_u4.data['total_chargeable_amount'], 1500.00)
        self.assertEqual(res_u4.data['net_payable_amt'], 1500.00)
        self.assertEqual(len(res_u4.data['chargeable_items']), 1)
        self.assertEqual(res_u4.data['chargeable_items'][0]['voucher_no'], 'VCH-U4-001')
        self.assertEqual(res_u4.data['production_unit']['name'], 'Unit 4')

        # Query Unit 2 statement
        res_u2 = self.client.get(self.bill_url, {
            'contractor': str(self.contractor.id),
            'production_unit': str(self.unit2.id),
            'month': 'Jul-26'
        })
        self.assertEqual(res_u2.status_code, status.HTTP_200_OK)
        self.assertEqual(res_u2.data['total_chargeable_amount'], 1000.00)
        self.assertEqual(res_u2.data['net_payable_amt'], 1000.00)
        self.assertEqual(len(res_u2.data['chargeable_items']), 1)
        self.assertEqual(res_u2.data['chargeable_items'][0]['voucher_no'], 'VCH-U2-001')
        self.assertEqual(res_u2.data['production_unit']['name'], 'Unit 2')

    def test_chargeable_included_and_non_chargeable_excluded(self):
        """Chargeable issues contribute to deduction; non-chargeable issues do not."""
        # Chargeable issue: 4 pcs @ 150 = 600.00
        StoreDailyIssue.objects.create(
            voucher_no='VCH-CHG-01',
            contractor=self.contractor,
            item=self.item1,
            qty=Decimal('4.000'),
            unit='pcs',
            rate=Decimal('150.00'),
            status=StoreItemStatus.CHARGE,
            production_unit=self.unit4,
            month_year='Jul-26',
            issued_by=self.admin
        )
        # Non-chargeable issue (free consumables): 2 ltr @ 200 = 400.00
        StoreDailyIssue.objects.create(
            voucher_no='VCH-NCH-01',
            contractor=self.contractor,
            item=self.item2,
            qty=Decimal('2.000'),
            unit='ltr',
            rate=Decimal('200.00'),
            status=StoreItemStatus.NON_CHARGE,
            production_unit=self.unit4,
            month_year='Jul-26',
            issued_by=self.admin
        )

        res = self.client.get(self.bill_url, {
            'contractor': str(self.contractor.id),
            'production_unit': str(self.unit4.id),
            'month': 'Jul-26'
        })
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        # Deduction should be exactly 600.00, NOT 1000.00
        self.assertEqual(res.data['total_chargeable_amount'], 600.00)
        self.assertEqual(res.data['net_payable_amt'], 600.00)
        self.assertEqual(res.data['total_non_chargeable_amount'], 400.00)
        self.assertEqual(res.data['chargeable_entries_count'], 1)
        self.assertEqual(res.data['non_chargeable_entries_count'], 1)
        self.assertEqual(res.data['total_issue_entries_count'], 2)

    def test_response_fields_contract_and_aliases(self):
        """API provides all canonical and legacy aliased fields with matching values."""
        StoreDailyIssue.objects.create(
            voucher_no='VCH-U4-ALIASES',
            contractor=self.contractor,
            item=self.item1,
            qty=Decimal('3.000'),
            unit='pcs',
            rate=Decimal('150.00'),
            status=StoreItemStatus.CHARGE,
            production_unit=self.unit4,
            month_year='Jul-26',
            issued_by=self.admin
        )

        res = self.client.get(self.bill_url, {
            'contractor': str(self.contractor.id),
            'production_unit': str(self.unit4.id),
            'month': 'Jul-26'
        })
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        data = res.data

        # Value agreement across all field aliases
        self.assertEqual(data['total_chargeable_amount'], 450.00)
        self.assertEqual(data['total_chargeable_amt'], 450.00)
        self.assertEqual(data['net_payable_amt'], 450.00)
        self.assertEqual(data['net_payable_amount'], 450.00)

        # Zero returns contract
        self.assertEqual(data['total_returned_amt'], 0.0)
        self.assertEqual(data['total_returned_amount'], 0.0)
        self.assertEqual(data['returned_items'], [])

        # Factory metadata
        self.assertEqual(data['production_unit']['id'], str(self.unit4.id))
        self.assertEqual(data['production_unit']['name'], 'Unit 4')
        self.assertEqual(data['production_unit']['unit_code'], 'U4_TEST')

    def test_unassigned_issues_detected_and_excluded_from_factory_bill(self):
        """Issues without production_unit are counted as unassigned and excluded from factory bills."""
        # Assigned to Unit 4: 5 pcs @ 150 = 750.00
        StoreDailyIssue.objects.create(
            voucher_no='VCH-ASSIGNED-01',
            contractor=self.contractor,
            item=self.item1,
            qty=Decimal('5.000'),
            unit='pcs',
            rate=Decimal('150.00'),
            status=StoreItemStatus.CHARGE,
            production_unit=self.unit4,
            month_year='Jul-26',
            issued_by=self.admin
        )

        # Unassigned issue: 3 pcs @ 150 = 450.00 (production_unit is None)
        StoreDailyIssue.objects.create(
            voucher_no='VCH-UNASSIGNED-01',
            contractor=self.contractor,
            item=self.item1,
            qty=Decimal('3.000'),
            unit='pcs',
            rate=Decimal('150.00'),
            status=StoreItemStatus.CHARGE,
            production_unit=None,
            month_year='Jul-26',
            issued_by=self.admin
        )

        res = self.client.get(self.bill_url, {
            'contractor': str(self.contractor.id),
            'production_unit': str(self.unit4.id),
            'month': 'Jul-26'
        })
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        # Bill must strictly total 750.00, NOT 1200.00
        self.assertEqual(res.data['total_chargeable_amount'], 750.00)
        self.assertEqual(res.data['net_payable_amt'], 750.00)
        self.assertEqual(len(res.data['chargeable_items']), 1)
        self.assertEqual(res.data['chargeable_items'][0]['voucher_no'], 'VCH-ASSIGNED-01')

        # Unassigned transaction is flagged
        self.assertEqual(res.data['unassigned_issues_count'], 1)

    def test_historical_material_returns_do_not_create_credits(self):
        """Historical StoreMaterialReturn entries do not deduct from contractor material bill."""
        # Issue: 10 pcs @ 150 = 1,500.00
        StoreDailyIssue.objects.create(
            voucher_no='VCH-ISSUE-RET-TEST',
            contractor=self.contractor,
            item=self.item1,
            qty=Decimal('10.000'),
            unit='pcs',
            rate=Decimal('150.00'),
            status=StoreItemStatus.CHARGE,
            production_unit=self.unit4,
            month_year='Jul-26',
            issued_by=self.admin
        )

        # Historical return entry: 2 pcs @ 150 = 300.00
        StoreMaterialReturn.objects.create(
            voucher_no='RET-HISTORICAL-01',
            contractor=self.contractor,
            item=self.item1,
            qty=Decimal('2.000'),
            unit='pcs',
            rate=Decimal('150.00'),
            status=StoreItemStatus.CHARGE,
            production_unit=self.unit4,
            month_year='Jul-26',
            returned_by=self.admin
        )

        res = self.client.get(self.bill_url, {
            'contractor': str(self.contractor.id),
            'production_unit': str(self.unit4.id),
            'month': 'Jul-26'
        })
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        # Net payable remains 1,500.00 (Returns are NOT deducted)
        self.assertEqual(res.data['total_chargeable_amount'], 1500.00)
        self.assertEqual(res.data['net_payable_amt'], 1500.00)
        self.assertEqual(res.data['total_returned_amt'], 0.0)

    def test_all_time_statement_support(self):
        """Month parameter supports 'all' for cross-month historical statements."""
        StoreDailyIssue.objects.create(
            voucher_no='VCH-M1',
            contractor=self.contractor,
            item=self.item1,
            qty=Decimal('2.000'),
            unit='pcs',
            rate=Decimal('150.00'),
            status=StoreItemStatus.CHARGE,
            production_unit=self.unit4,
            month_year='Jul-26',
            issued_by=self.admin
        )
        StoreDailyIssue.objects.create(
            voucher_no='VCH-M2',
            contractor=self.contractor,
            item=self.item1,
            qty=Decimal('3.000'),
            unit='pcs',
            rate=Decimal('150.00'),
            status=StoreItemStatus.CHARGE,
            production_unit=self.unit4,
            month_year='Aug-26',
            issued_by=self.admin
        )

        res = self.client.get(self.bill_url, {
            'contractor': str(self.contractor.id),
            'production_unit': str(self.unit4.id),
            'month': 'all'
        })
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res.data['total_chargeable_amount'], 750.00)
        self.assertEqual(res.data['month_year'], 'All Time')
        self.assertEqual(len(res.data['chargeable_items']), 2)

    def test_existing_store_endpoints_unbroken(self):
        """Related store endpoints continue to operate as expected."""
        res_issues = self.client.get(reverse('store-daily-issue-list'))
        self.assertEqual(res_issues.status_code, status.HTTP_200_OK)

        res_units = self.client.get(reverse('production-unit-list'))
        self.assertEqual(res_units.status_code, status.HTTP_200_OK)


class ContractorMaterialReturnRetirementTests(APITestCase):
    """
    Automated test suite verifying the complete retirement of contractor material returns:
    1. New return creation is rejected with HTTP 405.
    2. Historical returns cannot be updated or deleted via API.
    3. Historical returns remain readable for auditing.
    4. Excel import of material returns is rejected with clear validation error.
    5. Template downloads and exports for material returns are rejected.
    6. Standard 4-in-1 Excel import/export remains intact without returns.
    7. Inventory stock balance formula preserves historical return credits.
    8. Supplier debit-notes and procurement returns are completely unaffected.
    """

    def setUp(self):
        self.admin = User.objects.create_user(
            username='admin_return_audit',
            password='testpassword123',
            role=RoleChoices.ADMIN
        )
        self.client.force_authenticate(user=self.admin)

        self.unit1 = ProductionUnit.objects.create(name='Factory 1', unit_code='F1')
        self.contractor = User.objects.create_user(
            username='contractor_test',
            first_name='Dev',
            last_name='Carpenter',
            role=RoleChoices.CONTRACTOR,
            production_unit=self.unit1
        )
        self.cat = StoreItemCategory.objects.create(name='Hardware', code='HDW')
        self.item = StoreItem.objects.create(
            item_code='IT-SCRW-01',
            item_name='Wood Screw 2 inch',
            category=self.cat,
            unit='pcs',
            base_rate=Decimal('2.50'),
            current_rate=Decimal('2.50')
        )
        self.supplier = Supplier.objects.create(name='Hardware Mart')

    def test_new_contractor_return_post_rejected(self):
        """POST /api/store/material-returns/ must be rejected with HTTP 405 Method Not Allowed."""
        url = reverse('store-material-return-list')
        payload = {
            'voucher_no': 'RET-2026-999',
            'return_date': '2026-08-10',
            'contractor': self.contractor.id,
            'item': self.item.id,
            'qty': '50.00',
            'unit': 'pcs',
            'rate': '2.50',
            'production_unit': self.unit1.id,
            'remark': 'Attempt to return extra screws'
        }
        res = self.client.post(url, payload, format='json')
        self.assertEqual(res.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)
        self.assertIn('retired', res.data.get('detail', '').lower())
        self.assertEqual(StoreMaterialReturn.objects.count(), 0)

    def test_historical_return_update_rejected(self):
        """PUT / PATCH /api/store/material-returns/<id>/ must be rejected with HTTP 405."""
        historical_return = StoreMaterialReturn.objects.create(
            voucher_no='RET-HIST-001',
            return_date='2026-05-15',
            month_year='May-26',
            contractor=self.contractor,
            item=self.item,
            qty=Decimal('20.00'),
            unit='pcs',
            rate=Decimal('2.50'),
            total_amount=Decimal('50.00'),
            production_unit=self.unit1,
            remark='Historical return record',
            returned_by=self.admin
        )
        detail_url = reverse('store-material-return-detail', args=[historical_return.id])

        put_res = self.client.put(detail_url, {'qty': '100.00', 'remark': 'Tampered'}, format='json')
        self.assertEqual(put_res.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)
        self.assertIn('historical', put_res.data.get('detail', '').lower())

        patch_res = self.client.patch(detail_url, {'qty': '100.00'}, format='json')
        self.assertEqual(patch_res.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)

        historical_return.refresh_from_db()
        self.assertEqual(historical_return.qty, Decimal('20.00'))
        self.assertEqual(historical_return.remark, 'Historical return record')

    def test_historical_return_delete_rejected(self):
        """DELETE /api/store/material-returns/<id>/ must be rejected with HTTP 405."""
        historical_return = StoreMaterialReturn.objects.create(
            voucher_no='RET-HIST-002',
            return_date='2026-05-16',
            month_year='May-26',
            contractor=self.contractor,
            item=self.item,
            qty=Decimal('15.00'),
            unit='pcs',
            rate=Decimal('2.50'),
            total_amount=Decimal('37.50'),
            production_unit=self.unit1,
            returned_by=self.admin
        )
        detail_url = reverse('store-material-return-detail', args=[historical_return.id])

        del_res = self.client.delete(detail_url)
        self.assertEqual(del_res.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)
        self.assertTrue(StoreMaterialReturn.objects.filter(id=historical_return.id).exists())

    def test_historical_returns_readable_for_auditing(self):
        """GET /api/store/material-returns/ and GET /<id>/ must succeed for auditing."""
        historical_return = StoreMaterialReturn.objects.create(
            voucher_no='RET-HIST-003',
            return_date='2026-05-17',
            month_year='May-26',
            contractor=self.contractor,
            item=self.item,
            qty=Decimal('10.00'),
            unit='pcs',
            rate=Decimal('2.50'),
            total_amount=Decimal('25.00'),
            production_unit=self.unit1,
            remark='Audit check',
            returned_by=self.admin
        )

        list_res = self.client.get(reverse('store-material-return-list'))
        self.assertEqual(list_res.status_code, status.HTTP_200_OK)
        results = list_res.data.get('results', list_res.data)
        self.assertTrue(any(r['voucher_no'] == 'RET-HIST-003' for r in results))

        detail_res = self.client.get(reverse('store-material-return-detail', args=[historical_return.id]))
        self.assertEqual(detail_res.status_code, status.HTTP_200_OK)
        self.assertEqual(detail_res.data['voucher_no'], 'RET-HIST-003')

    def test_contractor_return_excel_import_rejected(self):
        """POST /api/store/import-excel/ with import_type='material_return' must be rejected with HTTP 400."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Material Returns"
        ws.append(["Voucher", "Date", "Item", "Contractor", "Qty", "Rate"])
        ws.append(["RET-TEST-01", "2026-08-01", self.item.item_code, self.contractor.username, 10, 2.50])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        excel_file = SimpleUploadedFile(
            'test_returns.xlsx',
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        res = self.client.post(
            reverse('store-import-excel'),
            {'file': excel_file, 'import_type': 'material_return'},
            format='multipart'
        )
        self.assertEqual(res.status_code, status.HTTP_400_BAD_REQUEST)
        error_msg = res.data.get('error') or res.data.get('detail') or ''
        self.assertIn('retired', error_msg.lower())

    def test_material_return_template_and_export_rejected(self):
        """Template download and selected export for material returns must return HTTP 400."""
        tmpl_res = self.client.get(reverse('store-export-template'), {'type': 'material_return'})
        self.assertEqual(tmpl_res.status_code, status.HTTP_400_BAD_REQUEST)

        export_res = self.client.post(reverse('store-export-selected'), {'module': 'material_return'}, format='json')
        self.assertEqual(export_res.status_code, status.HTTP_400_BAD_REQUEST)

        bulk_del_res = self.client.post(
            reverse('store-bulk-delete'),
            {'module': 'material_return', 'selected_ids': [1, 2, 3]},
            format='json'
        )
        self.assertEqual(bulk_del_res.status_code, status.HTTP_400_BAD_REQUEST)

    def test_standard_4in1_template_and_export_intact(self):
        """Unified 4-in-1 template download remains intact and does NOT include returns sheet."""
        tmpl_res = self.client.get(reverse('store-export-template'), {'type': 'unified_master'})
        self.assertEqual(tmpl_res.status_code, status.HTTP_200_OK)

        wb = openpyxl.load_workbook(io.BytesIO(tmpl_res.content))
        sheet_names = wb.sheetnames
        self.assertIn('Item Master', sheet_names)
        self.assertIn('Contractors-Supplier', sheet_names)
        self.assertIn('Material In', sheet_names)
        self.assertIn('Daily Issue Entry', sheet_names)
        self.assertNotIn('Material Returns', sheet_names)
        self.assertNotIn('Returns', sheet_names)

        export_items_res = self.client.post(reverse('store-export-selected'), {'module': 'items'}, format='json')
        self.assertEqual(export_items_res.status_code, status.HTTP_200_OK)

    def test_stock_balance_formula_preserves_historical_returns(self):
        """StoreItem balance calculations continue to include historical returns in inventory balance."""
        StoreMaterialIn.objects.create(
            voucher_no='INW-STK-001',
            inward_date='2026-05-01',
            supplier=self.supplier,
            item=self.item,
            qty=Decimal('100.00'),
            unit='pcs',
            bill_rate=Decimal('2.50'),
            total_amount=Decimal('250.00'),
            received_by=self.admin
        )
        StoreMaterialReturn.objects.create(
            voucher_no='RET-STK-001',
            return_date='2026-05-10',
            contractor=self.contractor,
            item=self.item,
            qty=Decimal('15.00'),
            unit='pcs',
            rate=Decimal('2.50'),
            total_amount=Decimal('37.50'),
            production_unit=self.unit1,
            returned_by=self.admin
        )
        StoreDailyIssue.objects.create(
            voucher_no='ISS-STK-001',
            issue_date='2026-05-12',
            contractor=self.contractor,
            item=self.item,
            qty=Decimal('25.00'),
            unit='pcs',
            rate=Decimal('2.50'),
            total_amount=Decimal('62.50'),
            production_unit=self.unit1,
            issued_by=self.admin
        )

        self.item.refresh_from_db()
        self.assertEqual(self.item.total_stock_qty, Decimal('100.00'))
        self.assertEqual(self.item.total_returned_qty, Decimal('15.00'))
        self.assertEqual(self.item.total_issued_qty, Decimal('25.00'))
        # Balance = Inward (100) + Returned (15) - Issued (25) = 90
        self.assertEqual(self.item.balance_stock_qty, Decimal('90.00'))

    def test_supplier_debit_notes_and_procurement_unaffected(self):
        """Supplier debit notes and tax invoice endpoints remain unaffected."""
        dn_res = self.client.get(reverse('supplier-debit-note-list'))
        self.assertEqual(dn_res.status_code, status.HTTP_200_OK)

        ti_res = self.client.get(reverse('supplier-tax-invoice-list'))
        self.assertEqual(ti_res.status_code, status.HTTP_200_OK)

        po_res = self.client.get(reverse('supplier-po-list'))
        self.assertEqual(po_res.status_code, status.HTTP_200_OK)


class BuyerPIExportTests(APITestCase):
    """
    Test suite for Buyer PI Excel export with multiple formats and item selection:
    1. General PI format with all items.
    2. General PI format with selective item filtering.
    3. EF PI format with all items.
    4. EF PI format with selective item filtering.
    """

    def setUp(self):
        from erp.models import Buyer, BuyerPI, BuyerPIItem, BuyerMaster, Sample, Finish, FinishCategory
        self.admin = User.objects.create_user(
            username='admin_pi_export',
            password='testpassword123',
            role=RoleChoices.ADMIN
        )
        self.client.force_authenticate(user=self.admin)

        self.buyer = Buyer.objects.create(
            name="Proceed Your Commerce S.L.",
            code="PYC01",
            address="Poligono Industrial La Florida Alicante Spain",
            phone="+34 678064958",
            email="purchases04@europefurni.com"
        )

        self.finish_walnut = Finish.objects.create(
            name="Walnut Finish",
            finish_code="W01",
            category=FinishCategory.WOOD,
            color="Walnut"
        )

        self.sample1 = Sample.objects.create(
            sample_id="SMP-NIGHT-01",
            product_name="Small 1 Drawer Nightstand",
            material="Mango",
            cbm=Decimal("0.1300"),
            usd=Decimal("47.50")
        )

        self.bm1 = BuyerMaster.objects.create(
            buyer=self.buyer,
            sample=self.sample1,
            style_no="2107-005-SMALL-Walnut",
            buyer_code="158591",
            product_name="Mesita de Noche de Madera de Mango Vanila",
            description="Slein Small 1 Drawer Nightstand",
            wood_type="Mango",
            wood_finish="Walnut",
            price_usd=Decimal("47.50"),
            cbm=Decimal("0.1300"),
            units=300,
            net_weight=Decimal("9.500"),
            gross_weight=Decimal("12.320"),
            box_length=Decimal("53"),
            box_breadth=Decimal("55"),
            box_height=Decimal("43")
        )

        self.bm2 = BuyerMaster.objects.create(
            buyer=self.buyer,
            style_no="2411-001-Natural-Headboard",
            buyer_code="178582",
            product_name="Cabecero para Cama de Madera de Mango",
            description="Slein Queen Headboard With Frame",
            wood_type="Mango",
            wood_finish="Natural",
            price_usd=Decimal("81.00"),
            cbm=Decimal("0.1500"),
            units=100,
            net_weight=Decimal("18.100"),
            gross_weight=Decimal("23.000"),
            box_length=Decimal("168"),
            box_breadth=Decimal("10"),
            box_height=Decimal("88")
        )

        self.pi = BuyerPI.objects.create(
            pi_no="060",
            buyer=self.buyer,
            pi_date="2026-08-12",
            ex_factory_date="2026-10-20",
            payment_terms="30% deposit and 70% Balance payment for Telex"
        )

        self.item1 = BuyerPIItem.objects.create(
            buyer_pi=self.pi,
            buyer_master=self.bm1,
            style_no=self.bm1.style_no,
            product_name=self.bm1.product_name,
            barcode="365771",
            buyer_no="158591",
            cbm=Decimal("0.1300"),
            price_usd=Decimal("47.50"),
            units=300,
            total_cbm=Decimal("39.0000"),
            total_amount=Decimal("14250.00")
        )

        self.item2 = BuyerPIItem.objects.create(
            buyer_pi=self.pi,
            buyer_master=self.bm2,
            style_no=self.bm2.style_no,
            product_name=self.bm2.product_name,
            barcode="438222",
            buyer_no="178582",
            cbm=Decimal("0.1500"),
            price_usd=Decimal("81.00"),
            units=100,
            total_cbm=Decimal("15.0000"),
            total_amount=Decimal("8100.00")
        )

    def test_export_general_format_all_items(self):
        url = reverse('buyer-pi-export-excel', args=[self.pi.id])
        res = self.client.get(url, {'export_format': 'general'})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", res['Content-Type'])
        self.assertIn("PI_060_General.xlsx", res['Content-Disposition'])

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        # Header + 2 items + totals row
        self.assertEqual(ws.cell(row=2, column=1).value, "S. No.")

    def test_export_general_format_selected_items(self):
        url = reverse('buyer-pi-export-excel', args=[self.pi.id])
        res = self.client.get(url, {'export_format': 'general', 'item_ids': str(self.item1.id)})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("PI_060_General_Selected.xlsx", res['Content-Disposition'])

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        # Total units should be 300 (only item1)
        self.assertEqual(ws.cell(row=4, column=14).value, 300)

    def test_export_ef_format_all_items(self):
        url = reverse('buyer-pi-export-excel', args=[self.pi.id])
        res = self.client.get(url, {'export_format': 'ef'})
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("PI_060_EF.xlsx", res['Content-Disposition'])

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws['A1'].value, "PINKCITY ENTERPRISES")
        self.assertEqual(ws.cell(row=13, column=1).value, "ITEM NO.")

    def test_export_ef_format_selected_items(self):
        url = reverse('buyer-pi-export-excel', args=[self.pi.id])
        res = self.client.post(url, {'export_format': 'ef', 'item_ids': [str(self.item2.id)]}, format='json')
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertIn("PI_060_EF_Selected.xlsx", res['Content-Disposition'])

        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        # Item row 14 should have item2 style
        self.assertEqual(ws.cell(row=14, column=1).value, self.bm2.style_no)
        # Total pcs on row 15 column 27 should be 100
        self.assertEqual(ws.cell(row=15, column=27).value, 100)

